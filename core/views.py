from datetime import date
from functools import wraps

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseBadRequest, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from .consent import CONSENT_TEXTS
from .forms import ConsentForm, ITEM_40_LABELS, RegisterForm, TestForm
from .models import Answer, TestSession, User
from .scoring import (
    ITEM_TEXTS,
    OPTION_LABELS,
    RESULTADOS_HEADERS,
    compute_results,
    item_score,
)


def admin_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect(f"{settings.LOGIN_URL}?next={request.path}")

        if not request.user.is_admin:
            messages.error(request, "Se requiere permiso de administrador.")
            return redirect("home")

        return view_func(request, *args, **kwargs)

    return _wrapped_view


def item_40_labels(answer_obj):
    if answer_obj is None:
        return []

    labels = []
    for option in answer_obj.selected_options or []:
        label = ITEM_40_LABELS.get(option, option)
        if option == "otros" and answer_obj.other_detail:
            label = f"Otros: {answer_obj.other_detail}"
        labels.append(label)
    return labels


def healthz(request):
    return HttpResponse("ok")


@login_required
def home(request):
    return render(request, "home.html")


def register(request):
    if request.user.is_authenticated:
        return redirect("home")

    if request.method == "POST":
        form = RegisterForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)

            # First registered user becomes administrator by default.
            if not User.objects.exists():
                user.is_admin = True
                user.is_staff = True
                user.is_superuser = True
                messages.success(
                    request,
                    "Primer usuario registrado. Se asignaron permisos de administrador.",
                )
            else:
                messages.success(request, "Cuenta creada correctamente.")

            user.save()
            login(request, user)
            return redirect("home")
    else:
        form = RegisterForm()

    return render(request, "register.html", {"form": form})


def _consent_session_key(mode: str) -> str:
    return f"fappren_consent_{mode}"


def _consent_url_name(mode: str) -> str:
    return "test_self" if mode == "self" else "test_administer"


def _test_form_url_name(mode: str) -> str:
    return "test_self_form" if mode == "self" else "test_administer_form"


def _test_started_session_key(mode: str) -> str:
    return f"fappren_test_started_{mode}"


@login_required
def informed_consent(request, mode: str):
    if mode not in CONSENT_TEXTS:
        return HttpResponseBadRequest("Modalidad inválida.")

    if request.method == "POST":
        form = ConsentForm(request.POST)
        if form.is_valid():
            request.session[_consent_session_key(mode)] = True
            return redirect(_test_form_url_name(mode))
    else:
        form = ConsentForm()

    return render(
        request,
        "consent.html",
        {
            "form": form,
            "consent_paragraphs": CONSENT_TEXTS[mode],
        },
    )


@login_required
def take_test(request, mode: str):
    if mode not in CONSENT_TEXTS:
        return HttpResponseBadRequest("Modalidad inválida.")

    if mode == "self":
        title = "Autoadministración del test FAPPREN VA"
        initial = {
            "mode": "self",
            "modality": "Autoadministración remota",
        }
    else:
        title = "Administración del test FAPPREN VA a un/a participante"
        initial = {
            "mode": "administered",
            "modality": "Asistida presencial",
        }

    if not request.session.get(_consent_session_key(mode)):
        messages.warning(
            request,
            "Debe aceptar el consentimiento informado para continuar.",
        )
        return redirect(_consent_url_name(mode))

    if request.method == "POST":
        form = TestForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data

            started_at = request.session.get(_test_started_session_key(mode))
            if started_at is None:
                time_minutes = None
            else:
                elapsed_seconds = max(0, timezone.now().timestamp() - started_at)
                time_minutes = max(1, round(elapsed_seconds / 60))

            session_obj = TestSession.objects.create(
                created_by=request.user,
                mode=mode,
                participant_code=data["participant_code"],
                age=data["age"],
                sex=data["sex"],
                education=data["education"],
                marital_status=data["marital_status"],
                occupational_status=data["occupational_status"],
                occupation_detail=data.get("occupation_detail", ""),
                nationality=data["nationality"],
                time_minutes=time_minutes,
                had_difficulty=data.get("had_difficulty", ""),
                difficulty_detail=data.get("difficulty_detail", ""),
                modality=data.get("modality", ""),
                evaluator_code=data.get("evaluator_code", ""),
                evaluator_read_aloud=data.get("evaluator_read_aloud", ""),
                received_help=data.get("received_help", ""),
                device=data.get("device", ""),
            )

            for item_number in range(1, 40):
                raw_value = data.get(f"item_{item_number}")
                if raw_value not in (None, ""):
                    Answer.objects.create(
                        session=session_obj,
                        item_number=item_number,
                        value=int(raw_value),
                    )

            item_40_selected = data.get("item_40") or []
            item_40_other = data.get("item_40_other", "")
            Answer.objects.create(
                session=session_obj,
                item_number=40,
                value=len(item_40_selected),
                selected_options=item_40_selected,
                other_detail=item_40_other if "otros" in item_40_selected else "",
            )

            request.session.pop(_consent_session_key(mode), None)
            request.session.pop(_test_started_session_key(mode), None)

            messages.success(request, "Respuestas guardadas correctamente.")
            return redirect("result_detail", pk=session_obj.pk)
    else:
        request.session[_test_started_session_key(mode)] = timezone.now().timestamp()
        form = TestForm(initial=initial)

    return render(
        request,
        "test_form.html",
        {
            "form": form,
            "title": title,
        },
    )


@login_required
def results_list(request):
    sessions = TestSession.objects.prefetch_related("answers").select_related("created_by")

    if not request.user.is_admin:
        sessions = sessions.filter(created_by=request.user)

    rows = []
    for session_obj in sessions:
        result = compute_results(session_obj.answers_dict())
        rows.append(
            {
                "obj": session_obj,
                "result": result,
            }
        )

    return render(request, "results_list.html", {"rows": rows})


@login_required
def result_detail(request, pk: int):
    session_obj = get_object_or_404(TestSession, pk=pk)

    if not request.user.is_admin and session_obj.created_by_id != request.user.id:
        return HttpResponseForbidden("No tiene permiso para ver este resultado.")

    answer_objects = {
        answer.item_number: answer for answer in session_obj.answers.all()
    }
    answers = {
        item_number: answer.value for item_number, answer in answer_objects.items()
    }
    result = compute_results(answers)

    answer_rows = []
    for item_number in range(1, 41):
        raw_value = answers.get(item_number)

        if item_number == 40:
            answer_obj = answer_objects.get(40)
            selected_labels = item_40_labels(answer_obj)
            raw_display = (
                ", ".join(selected_labels)
                if selected_labels
                else ("" if raw_value is None else str(raw_value))
            )
        elif raw_value is None:
            raw_display = ""
        else:
            raw_display = f"{raw_value} ({OPTION_LABELS.get(str(raw_value), '')})"

        answer_rows.append(
            {
                "item": item_number,
                "text": ITEM_TEXTS[item_number],
                "raw": raw_display,
                "score": item_score(item_number, raw_value),
            }
        )

    return render(
        request,
        "result_detail.html",
        {
            "test_session": session_obj,
            "result": result,
            "answer_rows": answer_rows,
        },
    )


@admin_required
def user_management(request):
    users = User.objects.all().order_by("username")
    return render(request, "user_management.html", {"users": users})


@admin_required
def toggle_admin(request, user_id: int):
    target_user = get_object_or_404(User, pk=user_id)

    if request.method != "POST":
        return redirect("user_management")

    if target_user == request.user:
        messages.error(request, "No puede modificar su propio permiso de administrador.")
        return redirect("user_management")

    if target_user.is_admin and User.objects.filter(is_admin=True).count() == 1:
        messages.error(request, "Debe existir al menos un administrador.")
        return redirect("user_management")

    target_user.is_admin = not target_user.is_admin
    target_user.is_staff = target_user.is_admin
    target_user.save()

    if target_user.is_admin:
        messages.success(request, f"{target_user.username} ahora es administrador.")
    else:
        messages.success(request, f"{target_user.username} ya no es administrador.")

    return redirect("user_management")


@admin_required
def export_results_xlsx(request):
    sessions = TestSession.objects.prefetch_related("answers").order_by("created_at")

    workbook = Workbook()

    worksheet_resultados = workbook.active
    worksheet_resultados.title = "Resultados"

    worksheet_resultados.append(RESULTADOS_HEADERS)

    for cell in worksheet_resultados[1]:
        cell.font = Font(bold=True)

    worksheet_resultados.freeze_panes = "A2"

    respuestas_headers = [
        "Código de participante",
        "Edad",
        "Sexo",
        "Nivel de escolaridad",
        "Estado civil",
        "Situación ocupacional",
        "Ocupación / profesión (detalle)",
        "Nacionalidad",
    ]

    respuestas_headers += [str(item_number) for item_number in range(1, 41)]
    respuestas_headers.append("Ítem 40 - Aspectos seleccionados")

    respuestas_headers += [
        "Tiempo de respuesta (minutos)",
        "¿Tuvo alguna dificultad para responder?",
        "Si tuvo dificultad: ¿cuál?",
        "Modalidad de aplicación",
        "Código del evaluador (solo asistida)",
        "Asistida: ¿el evaluador leyó los ítems en voz alta?",
        "Remota: ¿recibió ayuda de otra persona para responder?",
        "Soporte / dispositivo",
    ]

    worksheet_respuestas = workbook.create_sheet("Respuestas")
    worksheet_respuestas.append(respuestas_headers)

    for cell in worksheet_respuestas[1]:
        cell.font = Font(bold=True)

    worksheet_respuestas.freeze_panes = "A2"

    for session_obj in sessions:
        answer_objects = {
            answer.item_number: answer for answer in session_obj.answers.all()
        }
        answers = {
            item_number: answer.value
            for item_number, answer in answer_objects.items()
        }
        result = compute_results(answers)

        resultados_row = [session_obj.participant_code]

        for dimension_number in range(1, 11):
            score = result["dimension_scores"].get(dimension_number)
            resultados_row.append("" if score is None else score)

        resultados_row.append("" if result["total"] is None else result["total"])
        resultados_row.append(result["total_level"] or "")
        resultados_row.append(result["deficits_count"])
        resultados_row.append(result["areas"])

        for dimension_number in range(1, 11):
            resultados_row.append(
                result["dimension_levels"].get(dimension_number) or ""
            )

        worksheet_resultados.append(resultados_row)

        respuestas_row = [
            session_obj.participant_code,
            session_obj.age,
            session_obj.sex,
            session_obj.education,
            session_obj.marital_status,
            session_obj.occupational_status,
            session_obj.occupation_detail,
            session_obj.nationality,
        ]

        for item_number in range(1, 41):
            respuestas_row.append(answers.get(item_number, ""))

        item_40_answer = answer_objects.get(40)
        respuestas_row.append(", ".join(item_40_labels(item_40_answer)))

        # Application / modality metadata.
        respuestas_row.extend(
            [
                session_obj.time_minutes
                if session_obj.time_minutes is not None
                else "",
                session_obj.had_difficulty,
                session_obj.difficulty_detail,
                session_obj.modality,
                session_obj.evaluator_code,
                session_obj.evaluator_read_aloud,
                session_obj.received_help,
                session_obj.device,
            ]
        )

        worksheet_respuestas.append(respuestas_row)

    resultados_widths = [
        18,  # ID
        20,  # I
        20,  # II
        45,  # III
        35,  # IV
        35,  # V
        45,  # VI
        40,  # VII
        50,  # VIII
        35,  # IX
        40,  # X
        22,  # PUNTAJE TOTAL
        35,  # Nivel del total
        18,  # N° dimensiones en déficit
        55,  # Áreas de vulnerabilidad
        14,
        14,
        14,
        14,
        14,
        14,
        14,
        14,
        14,
        14,
    ]

    for index, width in enumerate(resultados_widths, start=1):
        worksheet_resultados.column_dimensions[
            worksheet_resultados.cell(row=1, column=index).column_letter
        ].width = width

    respuestas_widths = [
        22,  # Código participante
        8,   # Edad
        14,  # Sexo
        28,  # Escolaridad
        22,  # Estado civil
        32,  # Situación ocupacional
        30,  # Ocupación detalle
        18,  # Nacionalidad
    ]

    respuestas_widths += [6] * 40
    respuestas_widths.append(45)  # Ítem 40 - Aspectos seleccionados

    respuestas_widths += [
        18,
        25,
        35,
        28,
        22,
        35,
        38,
        20,
    ]

    for index, width in enumerate(respuestas_widths, start=1):
        worksheet_respuestas.column_dimensions[
            worksheet_respuestas.cell(row=1, column=index).column_letter
        ].width = width

    workbook.active = 0

    filename = f"FAPPREN_resultados_{date.today().isoformat()}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    workbook.save(response)
    return response